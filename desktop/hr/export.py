"""Export the HR records to an Excel workbook (.xlsx).

One sheet each for Employees, Attendance, Leave, Payroll, the Staff Report
and Licences. The file is a copy of the records at that moment: changes made
in Excel do not come back into the system.

Needs openpyxl:  pip install openpyxl
"""

import os
from datetime import date, datetime

from hr.constants import EMPLOYMENT_TYPES
from hr.database import APP_DIR, get_connection
from hr.reports import staff_summary

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:          # the rest of the app still works without it
    Workbook = None

OPENPYXL_MISSING = "Excel export needs openpyxl. Install it with:  pip install openpyxl"

HEADER_FILL = "C0501E"       # same orange as the app
TOTAL_FILL = "F7EFEA"
DATE_FORMAT = "yyyy-mm-dd"
MONEY_FORMAT = "#,##0.00"

# Employees sheet: (column in the database, heading, kind).
# kind is "text", "date", "money" or "number".
EMPLOYEE_COLUMNS = [
    ("employee_id", "Employee Code", "text"),
    ("first_name", "First Name", "text"),
    ("last_name", "Last Name", "text"),
    ("gender", "Gender", "text"),
    ("date_of_birth", "Date of Birth", "date"),
    ("phone", "Contact", "text"),
    ("emergency_contact", "Emergency Contact", "text"),
    ("email", "Email", "text"),
    ("address", "Address", "text"),
    ("ghana_card_no", "Ghana Card No", "text"),
    ("ssnit_no", "SSNIT No", "text"),
    ("qualification", "Qualification", "text"),
    ("employment_type", "Staff Status", "text"),
    ("company", "Company", "text"),
    ("service_status", "Service Status", "text"),
    ("department", "Department", "text"),
    ("position", "Designation", "text"),
    ("category", "Category", "text"),
    ("location", "Location", "text"),
    ("date_hired", "Date Hired", "date"),
    ("basic_salary", "Basic Salary", "money"),
    ("annual_leave", "Annual Leave Left", "number"),
    ("sick_leave", "Sick Leave Left", "number"),
    ("license_no", "License No", "text"),
    ("license_type", "License Type", "text"),
    ("license_first_renewal", "1st Renewal Date", "date"),
    ("license_second_renewal", "2nd Renewal Date", "date"),
    ("license_expiry", "Expiry Date", "date"),
]


def excel_available():
    return Workbook is not None


def default_filename():
    return f"HR Export {date.today().isoformat()}.xlsx"


def _as_date(value):
    """Turn "YYYY-MM-DD" text into a real date so Excel can sort and filter it."""
    if not value:
        return None
    try:
        return datetime.strptime(value, "%Y-%m-%d").date()
    except (TypeError, ValueError):
        return value          # keep anything unexpected as it was


def _add_sheet(workbook, title, headings, rows, kinds):
    """Write one formatted sheet: orange header, filters, frozen header row."""
    sheet = workbook.create_sheet(title)
    sheet.append(headings)
    for row in rows:
        sheet.append([
            _as_date(value) if kind == "date" else value
            for value, kind in zip(row, kinds)
        ])
        # Typed-in text starting with "=" must stay text, never become a
        # live Excel formula.
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.data_type = "s"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor=HEADER_FILL)
    for cell in sheet[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")

    for index, kind in enumerate(kinds, start=1):
        number_format = {"date": DATE_FORMAT, "money": MONEY_FORMAT}.get(kind)
        if number_format:
            for (cell,) in sheet.iter_rows(min_row=2, min_col=index, max_col=index):
                cell.number_format = number_format

    # Column widths from the longest value (within reason).
    for index, heading in enumerate(headings, start=1):
        longest = max(
            [len(str(heading))] + [
                10 if isinstance(cell.value, date) else len(str(cell.value or ""))
                for (cell,) in sheet.iter_rows(min_row=2, min_col=index, max_col=index)
            ]
        )
        sheet.column_dimensions[get_column_letter(index)].width = min(max(longest + 2, 8), 45)

    sheet.freeze_panes = "A2"
    if rows:
        sheet.auto_filter.ref = sheet.dimensions
    return sheet


def build_workbook():
    """Return an openpyxl Workbook holding all the HR records."""
    if Workbook is None:
        raise RuntimeError(OPENPYXL_MISSING)

    workbook = Workbook()
    workbook.remove(workbook.active)
    conn = get_connection()
    try:
        # ---- Employees ----
        employees = conn.execute(
            "SELECT * FROM employees WHERE active = 1 ORDER BY employee_id"
        ).fetchall()
        _add_sheet(
            workbook, "Employees",
            [heading for _, heading, _ in EMPLOYEE_COLUMNS],
            [[row[column] for column, _, _ in EMPLOYEE_COLUMNS] for row in employees],
            [kind for _, _, kind in EMPLOYEE_COLUMNS],
        )

        # ---- Attendance ----
        rows = conn.execute("""
            SELECT a.attendance_date, a.employee_id,
                   e.first_name || ' ' || e.last_name, a.shift,
                   a.check_in, a.check_out, a.status
            FROM attendance a JOIN employees e ON a.employee_id = e.employee_id
            ORDER BY a.attendance_date DESC, a.id DESC
        """).fetchall()
        _add_sheet(
            workbook, "Attendance",
            ["Date", "Employee Code", "Name", "Shift", "Check-in", "Check-out", "Status"],
            [list(r) for r in rows],
            ["date", "text", "text", "text", "text", "text", "text"],
        )

        # ---- Leave ----
        rows = conn.execute("""
            SELECT l.id, l.employee_id, e.first_name || ' ' || e.last_name,
                   l.leave_type, l.start_date, l.end_date, l.days, l.reason, l.status
            FROM leave_requests l JOIN employees e ON l.employee_id = e.employee_id
            ORDER BY l.id DESC
        """).fetchall()
        _add_sheet(
            workbook, "Leave",
            ["Request ID", "Employee Code", "Name", "Leave Type", "Start Date",
             "End Date", "Days", "Reason", "Status"],
            [list(r) for r in rows],
            ["number", "text", "text", "text", "date", "date", "number", "text", "text"],
        )

        # ---- Payroll ----
        rows = conn.execute("""
            SELECT p.employee_id, e.first_name || ' ' || e.last_name, p.month,
                   p.basic_salary, p.allowances, p.overtime, p.deductions, p.net_salary
            FROM payroll p JOIN employees e ON p.employee_id = e.employee_id
            ORDER BY p.month DESC, p.id DESC
        """).fetchall()
        _add_sheet(
            workbook, "Payroll",
            ["Employee Code", "Name", "Month", "Basic Salary", "Allowances",
             "Overtime", "Deductions", "Net Salary"],
            [list(r) for r in rows],
            ["text", "text", "text", "money", "money", "money", "money", "money"],
        )

        # ---- Licences ----
        today = date.today()
        licence_rows = []
        for row in employees:
            if not (row["license_no"] or row["license_expiry"]
                    or row["license_first_renewal"] or row["license_second_renewal"]):
                continue
            expiry = _as_date(row["license_expiry"])
            days_left = (expiry - today).days if isinstance(expiry, date) else None
            if days_left is None:
                status = ""
            elif days_left < 0:
                status = "Expired"
            elif days_left <= 30:
                status = "Due within 30 days"
            else:
                status = "OK"
            licence_rows.append([
                row["employee_id"], f"{row['first_name']} {row['last_name']}",
                row["position"], row["location"], row["license_no"],
                row["license_type"], row["license_first_renewal"],
                row["license_second_renewal"], row["license_expiry"],
                days_left, status,
            ])
        licence_rows.sort(key=lambda r: (r[9] is None, r[9] if r[9] is not None else 0))
        _add_sheet(
            workbook, "Licences",
            ["Employee Code", "Name", "Designation", "Location", "License No",
             "License Type", "1st Renewal Date", "2nd Renewal Date", "Expiry Date",
             "Days to Expiry", "Status"],
            licence_rows,
            ["text", "text", "text", "text", "text", "text", "date", "date", "date",
             "number", "text"],
        )
    finally:
        conn.close()

    # ---- Staff Report (put first: it's the overview) ----
    rows, totals = staff_summary()
    report = _add_sheet(
        workbook, "Staff Report",
        ["Location", *EMPLOYMENT_TYPES, "Total"],
        [list(r) for r in rows] + [["TOTAL", *(totals[k] for k in EMPLOYMENT_TYPES),
                                    totals["Total"]]],
        ["text"] + ["number"] * (len(EMPLOYMENT_TYPES) + 1),
    )
    report.auto_filter.ref = None
    for cell in report[report.max_row]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor=TOTAL_FILL)
    report.cell(row=report.max_row + 2, column=1,
                value=f"Exported {datetime.now():%Y-%m-%d %H:%M}").font = Font(italic=True)
    workbook.move_sheet(report, offset=-(len(workbook.sheetnames) - 1))
    workbook.active = 0

    return workbook


def export_to_excel(path):
    """Save the workbook to ``path`` and return the path."""
    build_workbook().save(path)
    return path


# ---------------------------------------------------------------------------
# Terminal menu
# ---------------------------------------------------------------------------


def print_export():
    """Save the workbook into an ``exports`` folder next to the app."""
    print("\n========== EXPORT TO EXCEL ==========")

    if not excel_available():
        print(OPENPYXL_MISSING)
        return

    folder = os.path.join(APP_DIR, "exports")
    os.makedirs(folder, exist_ok=True)
    path = os.path.join(folder, default_filename())

    try:
        export_to_excel(path)
    except PermissionError:
        print("Could not save. If that file is open in Excel, close it and try again.")
        return

    print("Saved:", path)
